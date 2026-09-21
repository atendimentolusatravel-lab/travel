#!/usr/bin/env python3
"""Gera cotacao-cristofer-reveillon.html a partir de cotacao-cristofer-reveillon.xlsx.

A planilha é a fonte da verdade. Este script lê as linhas da aba Cotação e os
parâmetros (câmbio, taxa de serviço) e escreve a página no padrão visual da
LusaTravel. Rode depois de cada atualização da planilha.
"""
import datetime as dt
import html
import os
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "cotacao-cristofer-reveillon.xlsx")
OUT = os.path.join(ROOT, "cotacao-cristofer-reveillon.html")

wb = load_workbook(XLSX)
p = wb["Parâmetros"]
s = wb["Cotação"]

params = {p.cell(r, 1).value: p.cell(r, 2).value for r in range(4, 11)}
EUR = float(params["Câmbio EUR → R$"])
USD = float(params["Câmbio USD → R$"])
CAMBIO_DATA = str(params["Data do câmbio"])
TAXA = float(params["Taxa de serviço LusaTravel (%)"])
PAX = int(params["Passageiros a partir de 05/01"])
RATES = {"BRL": 1.0, "EUR": EUR, "USD": USD}

rows = []
for r in range(5, 200):
    n = s.cell(r, 1).value
    if not isinstance(n, int):
        continue
    d = {
        "n": n, "cat": s.cell(r, 2).value, "item": s.cell(r, 3).value or "",
        "det": s.cell(r, 4).value or "", "data": s.cell(r, 5).value or "",
        "pax": s.cell(r, 6).value, "forn": s.cell(r, 7).value or "",
        "moeda": s.cell(r, 8).value or "EUR", "unit": s.cell(r, 9).value,
        "qtd": s.cell(r, 10).value or 1, "ref": s.cell(r, 12).value or "",
        "inc": (s.cell(r, 13).value or "Não") == "Sim",
        "status": s.cell(r, 15).value or "A cotar", "obs": s.cell(r, 16).value or "",
    }
    d["total_moeda"] = (float(d["unit"]) * float(d["qtd"])) if d["unit"] not in (None, "") else None
    d["total_brl"] = (d["total_moeda"] * RATES.get(d["moeda"], 1.0)) if (d["total_moeda"] is not None and d["inc"]) else 0.0
    rows.append(d)

CATS = ["Aéreo", "Hotel", "Trem", "Transfer", "Experiência", "Seguro/Docs", "Outros"]
CAT_LABEL = {"Aéreo": "Aéreo", "Hotel": "Hotelaria", "Trem": "Trens", "Transfer": "Transfers privativos",
             "Experiência": "Experiências e ingressos", "Seguro/Docs": "Seguro e documentos", "Outros": "Outros"}

subtotal = sum(d["total_brl"] for d in rows)
taxa_val = subtotal * TAXA
total = subtotal + taxa_val
cotados = sum(1 for d in rows if d["status"] not in ("A cotar", "Descartado"))
ativos = [d for d in rows if d["status"] != "Descartado"]


def brl(v):
    if v is None:
        return "—"
    return "R$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def money(v, moeda):
    if v is None:
        return "—"
    sym = {"BRL": "R$", "EUR": "€", "USD": "US$"}.get(moeda, moeda)
    return sym + " " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def esc(x):
    return html.escape(str(x)) if x is not None else ""


STATUS_CLASS = {"A cotar": "st-open", "Cotado": "st-quoted", "Aprovado pelo cliente": "st-ok",
                "Reservado": "st-booked", "Descartado": "st-off"}

now = dt.datetime.now(dt.timezone(dt.timedelta(hours=-3))).strftime("%d/%m/%Y %H:%M")

parts = []
parts.append(f"""<title>Cotação Cristofer Réveillon</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Poppins:wght@300;400;500;600;700&family=Montserrat:wght@400;500;600&family=Arimo:wght@700&display=swap">
<style>
:root{{
  --gold:#da8d00; --ink:#1a1510; --char:#333131; --mut:#9a8d80;
  --line:#e4ddd2; --paper:#faf7f2; --surf:#fff; --surf2:#f2ece3; --dest:#1a3a5c;
  --ok:#2d6a4f; --okbg:#e3f1ea; --book:#1a3a5c; --bookbg:#e4ecf5; --quo:#da8d00; --quobg:#fbf1dc;
  --sans:'Poppins',system-ui,-apple-system,'Segoe UI',sans-serif; --con:'Montserrat',var(--sans); --lbl:'Arimo',var(--sans);
}}
@media (prefers-color-scheme:dark){{ :root:not([data-theme="light"]){{
  --ink:#f0e8de; --char:#d9cec1; --mut:#8f8173; --line:#3c342a; --paper:#141109; --surf:#1e1a13; --surf2:#28231b;
  --gold:#e8a210; --dest:#3d78b8; --ok:#7fd1a5; --okbg:#1d3328; --book:#8fb8e8; --bookbg:#1c2a3a; --quo:#e8a210; --quobg:#3a2c10; }} }}
:root[data-theme="dark"]{{
  --ink:#f0e8de; --char:#d9cec1; --mut:#8f8173; --line:#3c342a; --paper:#141109; --surf:#1e1a13; --surf2:#28231b;
  --gold:#e8a210; --dest:#3d78b8; --ok:#7fd1a5; --okbg:#1d3328; --book:#8fb8e8; --bookbg:#1c2a3a; --quo:#e8a210; --quobg:#3a2c10; }}
*,*::before,*::after{{ box-sizing:border-box; margin:0; padding:0; }}
body{{ background:var(--paper); color:var(--char); font-family:var(--sans); font-size:14px; line-height:1.55; -webkit-font-smoothing:antialiased; }}
.wrap{{ max-width:1180px; margin:0 auto; padding:0 20px; }}
.mh{{ background:#1a1510; padding:40px 0 32px; position:relative; }}
.mh::after{{ content:''; position:absolute; left:0; right:0; bottom:0; height:3px; background:var(--gold); }}
.mh-logo{{ font-family:var(--con); font-weight:600; font-size:12px; letter-spacing:.34em; text-transform:uppercase; color:var(--gold); }}
.mh-kick{{ margin-top:18px; font-size:10px; font-weight:600; letter-spacing:.24em; text-transform:uppercase; color:rgba(240,232,222,.5); }}
.mh-name{{ margin-top:6px; font-weight:700; font-size:clamp(26px,4.5vw,38px); line-height:1.06; color:#fff; text-transform:uppercase; text-wrap:balance; }}
.mh-dest{{ margin-top:6px; font-weight:600; font-size:16px; color:var(--gold); }}
.mh-stats{{ margin-top:26px; display:grid; grid-template-columns:repeat(auto-fit,minmax(150px,1fr)); gap:0; border-top:1px solid rgba(240,232,222,.14); }}
.mh-stat{{ padding:16px 20px 2px 0; border-right:1px solid rgba(240,232,222,.12); margin-right:20px; }}
.mh-stat:last-child{{ border-right:0; }}
.mh-stat .v{{ font-weight:700; font-size:22px; color:#fff; line-height:1.1; font-variant-numeric:tabular-nums; }}
.mh-stat .v.g{{ color:var(--gold); }}
.mh-stat .k{{ margin-top:6px; font-size:9.5px; font-weight:500; letter-spacing:.18em; text-transform:uppercase; color:rgba(240,232,222,.45); }}
.sec{{ padding:34px 0 8px; }}
.eyebrow{{ font-size:10px; font-weight:600; letter-spacing:.22em; text-transform:uppercase; color:var(--gold); margin-bottom:6px; }}
.h2{{ font-weight:700; font-size:20px; color:var(--ink); line-height:1.2; display:flex; align-items:baseline; gap:12px; flex-wrap:wrap; }}
.h2 .sum{{ font-weight:500; font-size:13px; color:var(--mut); font-variant-numeric:tabular-nums; }}
.tw{{ overflow-x:auto; margin-top:14px; border:1px solid var(--line); background:var(--surf); }}
table{{ border-collapse:collapse; width:100%; min-width:900px; }}
th{{ font-family:var(--lbl); font-weight:700; font-size:10px; letter-spacing:.12em; text-transform:uppercase; color:var(--mut); background:var(--surf2); padding:10px 12px; text-align:left; border-bottom:1px solid var(--line); white-space:nowrap; }}
td{{ padding:11px 12px; border-bottom:1px solid var(--line); vertical-align:top; font-size:13px; }}
tr:last-child td{{ border-bottom:0; }}
td.n{{ color:var(--mut); font-variant-numeric:tabular-nums; width:34px; }}
td.item b{{ display:block; font-weight:600; color:var(--ink); font-size:13.5px; }}
td.item .det{{ display:block; margin-top:3px; font-weight:300; font-size:12px; color:var(--char); max-width:46ch; }}
td.item .ref{{ display:block; margin-top:5px; font-size:11px; color:var(--mut); font-style:italic; }}
td.item .obs{{ display:block; margin-top:5px; font-size:11.5px; color:var(--char); border-left:2px solid var(--gold); padding-left:8px; }}
td.dt{{ white-space:nowrap; color:var(--char); font-size:12.5px; }}
td.num{{ text-align:right; white-space:nowrap; font-variant-numeric:tabular-nums; }}
td.num .u{{ display:block; font-size:11px; color:var(--mut); }}
td.brl{{ text-align:right; white-space:nowrap; font-weight:600; color:var(--ink); font-variant-numeric:tabular-nums; }}
td.brl.zero{{ color:var(--mut); font-weight:400; }}
.chip{{ display:inline-block; font-family:var(--lbl); font-weight:700; font-size:9.5px; letter-spacing:.1em; text-transform:uppercase; padding:4px 8px; border:1px solid currentColor; white-space:nowrap; }}
.st-open{{ color:var(--mut); }}
.st-quoted{{ color:var(--quo); background:var(--quobg); }}
.st-ok{{ color:var(--ok); background:var(--okbg); }}
.st-booked{{ color:var(--book); background:var(--bookbg); border-color:var(--book); }}
.st-off{{ color:var(--mut); text-decoration:line-through; }}
.opt{{ display:inline-block; font-family:var(--lbl); font-weight:700; font-size:9.5px; letter-spacing:.1em; text-transform:uppercase; padding:3px 7px; margin-bottom:4px; background:var(--dest); color:#fff; }}
.opt.out{{ background:transparent; color:var(--mut); border:1px solid var(--line); }}
tr.off td{{ opacity:.55; }}
.tot{{ margin:34px 0 0; border-top:3px solid var(--gold); background:var(--surf); }}
.tot-row{{ display:grid; grid-template-columns:1fr auto; gap:16px; padding:12px 16px; border-bottom:1px solid var(--line); font-size:14px; }}
.tot-row .v{{ font-variant-numeric:tabular-nums; font-weight:600; color:var(--ink); }}
.tot-row.grand{{ background:#1a1510; color:#fff; font-weight:600; font-size:15px; border-bottom:0; }}
.tot-row.grand .v{{ color:var(--gold); font-size:18px; font-weight:700; }}
.note{{ margin-top:12px; font-size:12px; color:var(--mut); font-weight:300; }}
.note b{{ font-weight:600; color:var(--ink); }}
.legend{{ margin-top:18px; display:flex; flex-wrap:wrap; gap:8px 14px; font-size:11.5px; color:var(--mut); align-items:center; }}
.ft{{ margin-top:44px; background:#1a1510; padding:32px 0 28px; position:relative; }}
.ft::before{{ content:''; position:absolute; left:0; right:0; top:0; height:3px; background:var(--gold); }}
.ft-logo{{ font-family:var(--con); font-weight:600; font-size:12px; letter-spacing:.34em; text-transform:uppercase; color:var(--gold); }}
.ft-sig{{ margin-top:14px; display:flex; flex-wrap:wrap; gap:8px 28px; font-family:var(--con); font-size:13px; color:rgba(240,232,222,.72); }}
.ft-sig a{{ color:rgba(240,232,222,.72); text-decoration:none; border-bottom:1px solid rgba(218,141,0,.5); }}
.ft-addr{{ margin-top:12px; font-family:var(--con); font-size:11.5px; color:rgba(240,232,222,.38); }}
:focus-visible{{ outline:2px solid var(--gold); outline-offset:3px; }}
</style>

<header class="mh"><div class="wrap">
  <div class="mh-logo">LusaTravel</div>
  <div class="mh-kick">Cotação em andamento</div>
  <h1 class="mh-name">Cristofer</h1>
  <div class="mh-dest">Réveillon na Europa Central · 29/12/2026 a 12/01/2027 · 5 viajantes</div>
  <div class="mh-stats">
    <div class="mh-stat"><div class="v g">{brl(total)}</div><div class="k">Total cotado até agora</div></div>
    <div class="mh-stat"><div class="v">{cotados} de {len(rows)}</div><div class="k">Itens cotados</div></div>
    <div class="mh-stat"><div class="v">{brl(total / PAX)}</div><div class="k">Por pessoa</div></div>
    <div class="mh-stat"><div class="v">€ = {f"{EUR:.2f}".replace(".", ",")}</div><div class="k">Câmbio de {esc(CAMBIO_DATA)}</div></div>
    <div class="mh-stat"><div class="v">{now[:5]}</div><div class="k">Atualizado às {now[11:]}</div></div>
  </div>
</div></header>
""")

for cat in CATS:
    cr = [d for d in rows if d["cat"] == cat]
    if not cr:
        continue
    cat_total = sum(d["total_brl"] for d in cr)
    cat_cot = sum(1 for d in cr if d["status"] not in ("A cotar", "Descartado"))
    parts.append(f"""<section class="sec"><div class="wrap">
  <div class="eyebrow">{esc(cat)}</div>
  <h2 class="h2">{esc(CAT_LABEL[cat])} <span class="sum">{cat_cot} de {len(cr)} cotados · {brl(cat_total)}</span></h2>
  <div class="tw"><table>
  <thead><tr><th>#</th><th>Item</th><th>Data</th><th>Pax</th><th>Fornecedor</th><th style="text-align:right">Valor</th><th style="text-align:right">Total R$</th><th>Status</th></tr></thead>
  <tbody>""")
    for d in cr:
        off = d["status"] == "Descartado"
        opt = ""
        if cat == "Hotel" and "Opção" in d["item"]:
            opt = f'<span class="opt{"" if d["inc"] else " out"}">{"No total" if d["inc"] else "Alternativa"}</span><br>'
        elif not d["inc"] and not off:
            opt = '<span class="opt out">Fora do total</span><br>'
        val = "—"
        if d["total_moeda"] is not None:
            val = f'{money(d["total_moeda"], d["moeda"])}<span class="u">{money(float(d["unit"]), d["moeda"])} × {d["qtd"]}</span>'
        ref = f'<span class="ref">Ref. {esc(d["ref"])}</span>' if d["ref"] else ""
        obs = f'<span class="obs">{esc(d["obs"])}</span>' if d["obs"] else ""
        tb = brl(d["total_brl"]) if d["total_brl"] else "—"
        parts.append(f"""<tr{' class="off"' if off else ''}>
  <td class="n">{d["n"]}</td>
  <td class="item">{opt}<b>{esc(d["item"])}</b><span class="det">{esc(d["det"])}</span>{ref}{obs}</td>
  <td class="dt">{esc(d["data"])}</td>
  <td class="num">{esc(d["pax"])}</td>
  <td>{esc(d["forn"]) or '<span style="color:var(--mut)">—</span>'}</td>
  <td class="num">{val}</td>
  <td class="brl{' zero' if not d['total_brl'] else ''}">{tb}</td>
  <td><span class="chip {STATUS_CLASS.get(d["status"], "st-open")}">{esc(d["status"])}</span></td>
</tr>""")
    parts.append("</tbody></table></div></div></section>")

parts.append(f"""<section class="sec"><div class="wrap">
  <div class="tot">
    <div class="tot-row"><span>Subtotal dos itens incluídos ({sum(1 for d in rows if d["total_brl"])} cotados)</span><span class="v">{brl(subtotal)}</span></div>
    <div class="tot-row"><span>Taxa de serviço LusaTravel ({int(TAXA*100)}%)</span><span class="v">{brl(taxa_val)}</span></div>
    <div class="tot-row grand"><span>Total para o cliente até agora</span><span class="v">{brl(total)}</span></div>
  </div>
  <p class="note"><b>*Nada reservado, apenas cotado.</b> Câmbio {esc(CAMBIO_DATA)} · Euro = {f"{EUR:.2f}".replace(".", ",")} · Dólar = {f"{USD:.2f}".replace(".", ",")}. Faltam {len(ativos) - sum(1 for d in ativos if d["status"] != "A cotar")} itens a cotar. A estimativa preliminar do roteiro era R$ 140.000 a 200.000.</p>
  <div class="legend">
    <span class="chip st-open">A cotar</span><span class="chip st-quoted">Cotado</span><span class="chip st-ok">Aprovado pelo cliente</span><span class="chip st-booked">Reservado</span><span class="chip st-off">Descartado</span>
    <span>·</span><span class="opt">No total</span><span class="opt out">Alternativa</span><span>Só uma opção de hotel por cidade entra no total.</span>
  </div>
</div></section>

<footer class="ft"><div class="wrap">
  <div class="ft-logo">LusaTravel</div>
  <div class="ft-sig"><a href="mailto:atendimentolusatravel@gmail.com">atendimentolusatravel@gmail.com</a><a href="https://wa.me/5541991896076">+55 41 99189-6076</a></div>
  <div class="ft-addr">Alameda Princesa Izabel, 1700 · Bigorrilho · Curitiba – PR</div>
</div></footer>
""")

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(parts))
print(f"ok: {OUT} · {len(rows)} itens · cotados {cotados} · total {brl(total)}")
