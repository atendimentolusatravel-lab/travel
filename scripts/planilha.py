"""Utilidades para editar cotacao-cristofer-reveillon.xlsx com segurança.

Uso:
    from planilha import abrir, achar, inserir_depois, renumerar, salvar
"""
import os
from copy import copy

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "cotacao-cristofer-reveillon.xlsx")
CATS = ("Aéreo", "Hotel", "Trem", "Transfer", "Experiência", "Seguro/Docs", "Outros")


def abrir():
    wb = load_workbook(XLSX)
    return wb, wb["Cotação"], wb["Resumo"]


def achar(s, prefixo):
    """Linha do primeiro item cujo nome (coluna C) começa com o prefixo."""
    for r in range(5, s.max_row + 1):
        if str(s.cell(r, 3).value or "").startswith(prefixo):
            return r
    raise KeyError(prefixo)


def inserir_depois(s, r, n=1):
    """Insere n linhas depois de r copiando o estilo de r. Devolve a primeira nova linha."""
    s.insert_rows(r + 1, amount=n)
    for rr in range(r + 1, r + 1 + n):
        for c in range(1, 17):
            src, dst = s.cell(r, c), s.cell(rr, c)
            dst.font = copy(src.font); dst.fill = copy(src.fill); dst.border = copy(src.border)
            dst.alignment = copy(src.alignment); dst.number_format = src.number_format
        s.cell(rr, 2).value = s.cell(r, 2).value
        s.cell(rr, 13).value = "Não"; s.cell(rr, 15).value = "A cotar"
    return r + 1


def remover(s, prefixo):
    s.delete_rows(achar(s, prefixo))


def renumerar(s, z):
    """Renumera a coluna A, reescreve as fórmulas de cada item, dos totais e do Resumo."""
    n = 0; last = None
    for r in range(5, s.max_row + 1):
        if s.cell(r, 2).value in CATS:
            n += 1; s.cell(r, 1).value = n; last = r
            s.cell(r, 11).value = f'=IF(I{r}="","",I{r}*J{r})'
            s.cell(r, 14).value = (f'=IF(OR(M{r}<>"Sim",I{r}=""),0,K{r}*IF(H{r}="BRL",1,'
                                   f'IF(H{r}="EUR",Parâmetros!$B$4,IF(H{r}="USD",Parâmetros!$B$5,1))))')
    R0, RL = 5, last
    rt = None
    for r in range(RL + 1, s.max_row + 1):
        v = str(s.cell(r, 3).value or "")
        if v.startswith("SUBTOTAL"):
            s.cell(r, 14).value = f"=SUM(N{R0}:N{RL})"; rt = r
        elif v.startswith("Taxa de serviço"):
            s.cell(r, 14).value = f"=N{rt}*Parâmetros!$B$10"
        elif v.startswith("TOTAL PARA"):
            s.cell(r, 14).value = f"=N{rt}+N{rt+1}"
    for dv in s.data_validations.dataValidation:
        cols = {rg.split(":")[0].rstrip("0123456789") for rg in str(dv.sqref).split()}
        dv.sqref = " ".join(f"{c}{R0}:{c}{RL}" for c in sorted(cols))
    s.auto_filter.ref = f"A4:P{RL}"
    for i in range(5, 12):
        a = f"A{i}"
        z.cell(i, 2).value = f'=COUNTIF(Cotação!$B${R0}:$B${RL},{a})'
        z.cell(i, 3).value = f'=COUNTIFS(Cotação!$B${R0}:$B${RL},{a},Cotação!$O${R0}:$O${RL},"A cotar")'
        z.cell(i, 4).value = f'=COUNTIFS(Cotação!$B${R0}:$B${RL},{a},Cotação!$O${R0}:$O${RL},"<>A cotar")'
        z.cell(i, 5).value = f'=SUMIF(Cotação!$B${R0}:$B${RL},{a},Cotação!$N${R0}:$N${RL})'
    return n


def salvar(wb):
    wb.save(XLSX)
