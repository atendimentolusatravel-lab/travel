#!/usr/bin/env python3
"""Gera a planilha do roteiro de voos de Luanna Alves Silva (fonte: caderno de vouchers JDFW5X / JDFX2R)."""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).parent
VERDE, CREME, LINHA = "404626", "FAF9F5", "E7E7E0"

COLS = ["Nº", "Data", "Dia", "Companhia", "Voo", "Operado por", "Origem", "Aeroporto de origem", "Partida",
        "Destino", "Aeroporto de destino", "Chegada", "Duração", "Aeronave", "Classe", "Tarifa",
        "Bagagem despachada", "Reserva", "Localizador cia.", "Bilhete", "Observação"]
VOOS = [
    [1, "02/12/2026", "qua", "Hahn Air", "H1 6425", "Safarilink Aviation", "WIL", "Nairóbi – Wilson Airport", "07:15",
     "ARK", "Arusha", "08:15", "1h00", "De Havilland Dash 8", "T", "Econômica", "15 kg", "JDFW5X", "BB41ZD",
     "169-5860070910", "Check-in no balcão da Safarilink; chegar 1h antes"],
    [2, "05/12/2026", "sáb", "Hahn Air", "H1 6427", "Safarilink Aviation", "ARK", "Arusha", "14:40",
     "WIL", "Nairóbi – Wilson Airport", "15:40", "1h00", "De Havilland Dash 8", "T", "Econômica", "15 kg", "JDFW5X",
     "BB41ZD", "169-5860070910", "Pouso no Wilson (aeroporto doméstico)"],
    [3, "06/12/2026", "dom", "Kenya Airways", "KQ 250", "Kenya Airways", "NBO", "Nairóbi – Jomo Kenyatta, Terminal 1A",
     "08:40", "SEZ", "Mahé, Seicheles – terminal internacional", "12:55", "cerca de 3h15", "Boeing 737-800", "N",
     "Econômica", "2 volumes", "JDFX2R", "BCAX3T", "169-5860070835",
     "Embarque no Jomo Kenyatta (não no Wilson); chegar 3h antes; Travel Authorization das Seicheles"],
]

wb = Workbook()
ws = wb.active
ws.title = "Voos"
fino = Side(style="thin", color=LINHA)
borda = Border(left=fino, right=fino, top=fino, bottom=fino)

ws["A1"] = "Roteiro de voos — Luanna Alves Silva"
ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=VERDE)
ws["A2"] = "Passageira: SILVA / LUANNA ALVES — SRA. · Quênia, Tanzânia e Seicheles · 02 a 06/12/2026 · bilhetes emitidos em 14/10/2026"
ws["A3"] = "Horários locais de cada aeroporto. As Seicheles estão 1 hora à frente de Nairóbi."
for c in ("A2", "A3"):
    ws[c].font = Font(name="Calibri", size=10, color="6F6F68")

H = 5
for j, nome in enumerate(COLS, 1):
    cel = ws.cell(row=H, column=j, value=nome)
    cel.font = Font(name="Calibri", bold=True, color="FFFFFF")
    cel.fill = PatternFill("solid", fgColor=VERDE)
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cel.border = borda
for i, v in enumerate(VOOS):
    for j, val in enumerate(v, 1):
        cel = ws.cell(row=H + 1 + i, column=j, value=val)
        cel.alignment = Alignment(vertical="top", wrap_text=True,
                                  horizontal="center" if j in (1, 3, 5, 7, 9, 10, 12, 15) else "left")
        cel.border = borda
        if i % 2:
            cel.fill = PatternFill("solid", fgColor=CREME)
        if j in (5, 9, 12):
            cel.font = Font(name="Calibri", bold=True)

LARG = [5, 12, 6, 14, 10, 18, 8, 30, 9, 9, 32, 9, 13, 20, 7, 11, 12, 10, 12, 16, 48]
for j, w in enumerate(LARG, 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[H].height = 32
ws.freeze_panes = ws.cell(row=H + 1, column=1)
ws.auto_filter.ref = f"A{H}:{get_column_letter(len(COLS))}{H + len(VOOS)}"

r = H + len(VOOS) + 2
ws.cell(row=r, column=1, value="Importante").font = Font(name="Calibri", bold=True, color=VERDE)
for k, t in enumerate([
    "Em 05/12 o pouso é no Wilson (WIL); em 06/12 o embarque é no Jomo Kenyatta (NBO), Terminal 1A. São aeroportos diferentes, a cerca de 20 km.",
    "Voos para Arusha (Hahn Air, operados pela Safarilink): check-in no balcão da Safarilink. Avião pequeno: 15 kg por pessoa, de preferência em mala macia.",
    "Seicheles: preencher a Travel Authorization em seychelles.govtas.com antes do embarque. Sem ela, a companhia não embarca.",
    "Acompanhante: DE PROSPERO AMUI / EDUARDO HENRIQUE — SR. · bilhetes 169-5860070911 (JDFW5X) e 169-5860070836 (JDFX2R).",
    "Emergência 24h LusaTravel: Claudia +55 41 99275-0065 · Victoria +55 41 99912-1919.",
], 1):
    ws.cell(row=r + k, column=1, value="• " + t).font = Font(name="Calibri", size=10)

ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

out = ROOT / "voos-luanna-alves-silva.xlsx"
wb.save(out)
print(out)
